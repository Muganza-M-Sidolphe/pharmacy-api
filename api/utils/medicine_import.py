"""Reading medicine files (CSV, XLSX, text PDF) and validating their rows for import."""

import calendar
import csv
import io
import re
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

MAX_FILE_SIZE = 10 * 1024 * 1024
MAX_ROWS = 5000
HEADER_SCAN_ROWS = 15

TARGET_FIELDS = [
    "medicine_name", "generic_name", "category", "strength", "dosage_form", "manufacturer",
    "barcode", "batch_number", "manufacturing_date", "expiry_date", "quantity", "cost_price",
    "selling_price", "supplier", "unit",
]
IGNORE = "ignore"

FIELD_LABELS = {
    "medicine_name": "Medicine name",
    "generic_name": "Generic name",
    "batch_number": "Batch number",
    "manufacturing_date": "Manufacturing date",
    "expiry_date": "Expiry date",
    "quantity": "Quantity",
    "cost_price": "Cost price",
    "selling_price": "Selling price",
}

# Exact header names (normalized) → field. Matches are reported with high confidence.
HEADER_ALIASES = {
    "medicine": "medicine_name", "medicinename": "medicine_name", "product": "medicine_name",
    "productname": "medicine_name", "drug": "medicine_name", "drugname": "medicine_name",
    "brand": "medicine_name", "brandname": "medicine_name", "itemname": "medicine_name",
    "name": "medicine_name", "item": "medicine_name",
    "generic": "generic_name", "genericname": "generic_name", "inn": "generic_name",
    "category": "category", "therapeuticclass": "category", "class": "category",
    "strength": "strength", "dose": "strength", "dosage": "strength",
    "dosageform": "dosage_form", "form": "dosage_form",
    "manufacturer": "manufacturer", "mfr": "manufacturer", "maker": "manufacturer",
    "barcode": "barcode", "upc": "barcode", "ean": "barcode",
    "batch": "batch_number", "batchno": "batch_number", "batchnumber": "batch_number",
    "lot": "batch_number", "lotno": "batch_number", "lotcode": "batch_number", "lotnumber": "batch_number",
    "mfgdate": "manufacturing_date", "manufacturingdate": "manufacturing_date",
    "manufacturedate": "manufacturing_date", "productiondate": "manufacturing_date",
    "expiry": "expiry_date", "expiration": "expiry_date", "expirationdate": "expiry_date",
    "expirydate": "expiry_date", "expdate": "expiry_date", "exp": "expiry_date",
    "quantity": "quantity", "qty": "quantity", "qoh": "quantity", "qtyonhand": "quantity",
    "stock": "quantity", "stockquantity": "quantity", "quantityinstock": "quantity",
    "cost": "cost_price", "costprice": "cost_price", "unitcost": "cost_price",
    "purchaseprice": "cost_price", "buyingprice": "cost_price",
    "sellingprice": "selling_price", "saleprice": "selling_price", "retailprice": "selling_price",
    "unitprice": "selling_price",
    "supplier": "supplier", "suppliername": "supplier", "vendor": "supplier",
    "unit": "unit", "uom": "unit", "unitofmeasure": "unit",
}

# Partial matches (substring of the normalized header), checked in order.
# Reported with medium confidence because they need a human glance.
HEADER_HINTS = [
    (("expir", "bestbefore", "useby"), "expiry_date"),
    (("mfg", "manufacturingdate", "manufacturedate", "productiondate"), "manufacturing_date"),
    (("batch", "lot"), "batch_number"),
    (("generic",), "generic_name"),
    (("barcode", "ean", "upc", "gtin"), "barcode"),
    (("selling", "retail", "saleprice", "mrp"), "selling_price"),
    (("cost", "purchase", "buying"), "cost_price"),
    (("price", "rate"), "selling_price"),
    (("qty", "quantity", "stock", "balance"), "quantity"),
    (("manufactur",), "manufacturer"),
    (("supplier", "vendor", "distributor"), "supplier"),
    (("strength",), "strength"),
    (("dosage", "form"), "dosage_form"),
    (("categor",), "category"),
    (("medicine", "drug", "product", "brand", "item", "name"), "medicine_name"),
    (("unit", "uom", "pack"), "unit"),
]


class ImportFileError(Exception):
    """The uploaded file cannot be read as a medicine list."""


def normalize_header(value):
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def normalize_name(value):
    return re.sub(r"\s+", " ", str(value or "").strip()).lower()


def suggest_field(header):
    """Return (target field, confidence) for a column header."""
    key = normalize_header(header)
    if not key:
        return IGNORE, "low"
    if key in HEADER_ALIASES:
        return HEADER_ALIASES[key], "high"
    for needles, field in HEADER_HINTS:
        if any(needle in key for needle in needles):
            return field, "medium"
    return IGNORE, "low"


# ---------------------------------------------------------------------------
# File reading
# ---------------------------------------------------------------------------

def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _read_csv(data):
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [[_cell_text(cell) for cell in row] for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ImportFileError("The Excel file could not be opened. Save it as .xlsx and try again.") from exc
    try:
        # Use the first sheet that has any data.
        for sheet in workbook.worksheets:
            rows = [[_cell_text(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
            if any(any(row) for row in rows):
                return rows
        return []
    finally:
        workbook.close()


def _pdf_lines(page):
    """Words grouped into lines of cells [(x0, x1, text)]; a gap wider than the text height starts a new cell."""
    lines = []
    for word in sorted(page.extract_words(), key=lambda w: (round(w["top"]), w["x0"])):
        if lines and abs(word["top"] - lines[-1]["top"]) <= 3:
            line = lines[-1]
        else:
            line = {"top": word["top"], "cells": []}
            lines.append(line)
        cells = line["cells"]
        height = word["bottom"] - word["top"]
        if cells and word["x0"] - cells[-1][1] <= height:
            x0, _, text = cells[-1]
            cells[-1] = (x0, word["x1"], f"{text} {word['text']}")
        else:
            cells.append((word["x0"], word["x1"], word["text"]))
    return [line["cells"] for line in lines]


def _align_pdf_lines(lines, anchor):
    """Place each cell under the anchor (header) column it overlaps, or the nearest one."""
    rows = []
    for cells in lines:
        row = [""] * len(anchor)
        for x0, x1, text in cells:
            def distance(column):
                a0, a1 = column
                overlap = min(x1, a1) - max(x0, a0)
                return -overlap if overlap > 0 else abs((x0 + x1) / 2 - (a0 + a1) / 2)
            index = min(range(len(anchor)), key=lambda i: distance(anchor[i]))
            row[index] = f"{row[index]} {text}".strip()
        rows.append(row)
    return rows


def _read_pdf(data):
    import pdfplumber

    rows = []
    has_text = False
    anchor = None
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables()
                if tables:  # ruled table
                    has_text = True
                    for table in tables:
                        rows.extend([[_cell_text(cell) for cell in row] for row in table])
                    continue
                lines = _pdf_lines(page)
                if not lines:
                    continue
                has_text = True
                # The header is the line naming the most known fields; reuse the last page's if none.
                scored = [
                    (sum(1 for _, _, text in cells if suggest_field(text)[0] != IGNORE), len(cells), cells)
                    for cells in lines[:HEADER_SCAN_ROWS]
                ]
                best = max(scored, key=lambda item: (item[0], item[1]))
                if best[0] or anchor is None:
                    anchor = [(x0, x1) for x0, x1, _ in best[2]]
                rows.extend(_align_pdf_lines(lines, anchor))
    except Exception as exc:
        raise ImportFileError("The PDF file could not be read.") from exc
    if not has_text and not rows:
        raise ImportFileError(
            "This PDF looks like a scanned image. Text recognition (OCR) is not available yet; "
            "export the list to Excel or CSV and upload that instead."
        )
    return rows


def read_rows(file_name, data):
    """Return (file type, list of rows as lists of strings)."""
    extension = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if extension == "csv":
        return "csv", _read_csv(data)
    if extension == "xlsx":
        return "xlsx", _read_xlsx(data)
    if extension == "pdf":
        return "pdf", _read_pdf(data)
    raise ImportFileError("Unsupported file type. Upload an Excel (.xlsx), CSV or PDF file.")


def extract_table(rows):
    """Find the header row and return (column names, [(row number, {column: value})])."""
    best_index, best_score = None, 0
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = sum(1 for cell in row if suggest_field(cell)[0] != IGNORE)
        if score > best_score:
            best_index, best_score = index, score
    if best_index is None:
        best_index = next((i for i, row in enumerate(rows) if sum(1 for cell in row if cell) >= 2), None)
    if best_index is None:
        raise ImportFileError("No header row was found. The file needs a row of column names.")

    header = rows[best_index]
    body = rows[best_index + 1:]
    width = max([len(header)] + [len(row) for row in body])

    columns, seen = [], {}
    for position in range(width):
        name = header[position] if position < len(header) else ""
        has_data = any(position < len(row) and row[position] for row in body)
        if not name and not has_data:
            columns.append(None)
            continue
        name = name or f"Column {position + 1}"
        seen[name] = seen.get(name, 0) + 1
        if seen[name] > 1:
            name = f"{name} ({seen[name]})"
        columns.append(name)

    records = []
    for offset, row in enumerate(body):
        if not any(row) or row[:len(header)] == header:
            continue  # blank line or a header repeated on a new PDF page
        values = {
            column: (row[position] if position < len(row) else "")
            for position, column in enumerate(columns)
            if column
        }
        records.append((best_index + offset + 2, values))
        if len(records) > MAX_ROWS:
            raise ImportFileError(f"The file has more than {MAX_ROWS} rows. Split it into smaller files.")
    return [column for column in columns if column], records


def detect_columns(columns):
    """Suggest a target field per column; each field is suggested for one column only.

    Exact header matches claim their field first, then partial matches, in column order.
    """
    detected = [{"name": name, "target_field": IGNORE, "confidence": "low"} for name in columns]
    suggestions = [suggest_field(name) for name in columns]
    taken = set()
    for wanted in ("high", "medium"):
        for item, (field, confidence) in zip(detected, suggestions):
            if confidence == wanted and field not in taken:
                item["target_field"], item["confidence"] = field, confidence
                taken.add(field)
    return detected


# ---------------------------------------------------------------------------
# Value parsing
# ---------------------------------------------------------------------------

def parse_decimal(value):
    """Parse '1,500 RWF' → Decimal('1500'). Returns None for blank, raises ValueError if invalid."""
    text = str(value or "").strip()
    if not text:
        return None
    cleaned = re.sub(r"[^0-9.\-]", "", text.replace(",", ""))
    if not cleaned or cleaned.count(".") > 1 or cleaned.count("-") > 1 or "-" in cleaned[1:]:
        raise ValueError(text)
    try:
        return Decimal(cleaned)
    except InvalidOperation as exc:
        raise ValueError(text) from exc


DATE_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y",
    "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y", "%m/%d/%Y",
]
MONTH_FORMATS = ["%m/%Y", "%m-%Y", "%Y-%m", "%Y/%m", "%b %Y", "%B %Y", "%b-%Y", "%b-%y", "%m/%y"]


def parse_date(value, month_end=False):
    """Parse common date spellings; month-only dates use the first or last day of the month.

    Returns None for blank, raises ValueError if the text is not a date.
    Day-first (dd/mm/yyyy) is assumed when a date is ambiguous.
    """
    text = str(value or "").strip()
    if not text:
        return None
    text = re.sub(r"[T ]00:00(:00)?$", "", text)
    if re.fullmatch(r"\d{5}(\.0+)?", text):
        serial = int(float(text))
        if 20000 <= serial <= 80000:  # Excel serial date (1954–2119)
            return date(1899, 12, 30) + timedelta(days=serial)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    for fmt in MONTH_FORMATS:
        try:
            parsed = datetime.strptime(text, fmt).date()
        except ValueError:
            continue
        if month_end:
            return parsed.replace(day=calendar.monthrange(parsed.year, parsed.month)[1])
        return parsed
    raise ValueError(text)


def money(value):
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


# ---------------------------------------------------------------------------
# Record evaluation
# ---------------------------------------------------------------------------

def generated_batch_number(session, record):
    # AUTO-YYYYMMDD-NNNNNNNN: import date, 4 digits from the session id, 4-digit row number.
    session_digits = int(session.id.hex[:6], 16) % 10000
    return f"AUTO-{session.created_at:%Y%m%d}-{session_digits:04d}{record.row_number % 10000:04d}"


def record_values(session, record):
    """Mapped text for every target field: correction first, else the first non-empty mapped column."""
    values = {field: "" for field in TARGET_FIELDS}
    for column, field in session.mapping.items():
        if field in values and not values[field]:
            values[field] = str(record.raw.get(column, "") or "").strip()
    for field, value in (record.overrides or {}).items():
        if field in values:
            values[field] = str(value if value is not None else "").strip()
    return values


def evaluate_records(session, records, existing_medicines, existing_batches, today=None):
    """Validate records in order. Returns a list of result dicts (see evaluate_record).

    existing_medicines: {normalized name: Medicine}
    existing_batches: set of (medicine id, normalized batch number)
    """
    today = today or date.today()
    seen_batches = {}
    return [
        evaluate_record(session, record, existing_medicines, existing_batches, seen_batches, today)
        for record in records
    ]


def evaluate_record(session, record, existing_medicines, existing_batches, seen_batches, today):
    values = record_values(session, record)
    errors, warnings = [], []
    parsed = {}

    name = values["medicine_name"]
    if not name:
        errors.append("Medicine name is missing.")

    # Quantity
    try:
        quantity = parse_decimal(values["quantity"])
    except ValueError:
        quantity = None
        errors.append(f"Quantity '{values['quantity']}' is not a number.")
    else:
        if quantity is None:
            errors.append("Quantity is missing.")
        elif quantity != quantity.to_integral_value():
            errors.append("Quantity must be a whole number.")
        elif quantity < 0:
            errors.append("Quantity cannot be negative.")
        else:
            parsed["quantity"] = int(quantity)
            values["quantity"] = str(int(quantity))
            if quantity == 0:
                warnings.append("Quantity is 0.")

    # Prices
    for field in ("cost_price", "selling_price"):
        try:
            amount = parse_decimal(values[field])
        except ValueError:
            errors.append(f"{FIELD_LABELS[field]} '{values[field]}' is not a number.")
            continue
        if amount is not None:
            if amount < 0:
                errors.append(f"{FIELD_LABELS[field]} cannot be negative.")
                continue
            parsed[field] = money(amount)
            values[field] = str(parsed[field])

    price_source = "file"
    if "selling_price" in (record.overrides or {}):
        price_source = "manual"
    if "selling_price" not in parsed and not any(
        e.startswith(FIELD_LABELS["selling_price"]) for e in errors
    ):
        if session.pricing_method == "markup" and session.markup_percent is not None:
            if "cost_price" in parsed:
                markup = Decimal(session.markup_percent)
                parsed["selling_price"] = money(parsed["cost_price"] * (1 + markup / 100))
                values["selling_price"] = str(parsed["selling_price"])
                price_source = "markup"
            else:
                errors.append("Cost price is needed to calculate the selling price with markup.")
        elif session.pricing_method == "manual":
            errors.append("Enter a selling price.")
        else:
            errors.append("Selling price is missing.")
    if "cost_price" not in parsed and not values["cost_price"]:
        warnings.append("Cost price is missing; it will be saved as 0.")
    if "cost_price" in parsed and "selling_price" in parsed and parsed["selling_price"] < parsed["cost_price"]:
        warnings.append("Selling price is lower than cost price.")

    # Dates
    for field, month_end in (("manufacturing_date", False), ("expiry_date", True)):
        try:
            parsed_date = parse_date(values[field], month_end=month_end)
        except ValueError:
            errors.append(f"{FIELD_LABELS[field]} '{values[field]}' is not a valid date.")
            continue
        if parsed_date:
            parsed[field] = parsed_date
            values[field] = parsed_date.isoformat()
    expiry = parsed.get("expiry_date")
    manufactured = parsed.get("manufacturing_date")
    if not values["expiry_date"]:
        warnings.append("No expiry date.")
    elif expiry and expiry < today:
        warnings.append("This batch has already expired.")
    if expiry and manufactured and manufactured > expiry:
        errors.append("Manufacturing date is after the expiry date.")

    # Batch number
    batch_generated = False
    batch_source = "manual" if "batch_number" in (record.overrides or {}) else "file"
    if not values["batch_number"]:
        if session.missing_batch_strategy == "generate":
            values["batch_number"] = generated_batch_number(session, record)
            batch_generated = True
            batch_source = "system"
        else:
            errors.append("Batch number is missing.")

    # Medicine and duplicate batch checks
    name_key = normalize_name(name)
    existing = existing_medicines.get(name_key) if name else None
    batch_key = normalize_name(values["batch_number"])
    if name and batch_key:
        if existing and (existing.id, batch_key) in existing_batches:
            errors.append(f"Batch {values['batch_number']} already exists in inventory for this medicine.")
        elif (name_key, batch_key) in seen_batches:
            errors.append(
                f"Batch {values['batch_number']} for this medicine is repeated (first on row {seen_batches[(name_key, batch_key)]})."
            )
        else:
            seen_batches[(name_key, batch_key)] = record.row_number

    is_valid = not errors
    return {
        "id": str(record.id),
        "row_number": record.row_number,
        "status": "error" if errors else ("warning" if warnings else "valid"),
        "is_valid": is_valid,
        "errors": errors,
        "warnings": warnings,
        "mapped_data": values,
        "batch_number_generated": batch_generated,
        "batch_number_source": batch_source,
        "selling_price_source": price_source,
        "medicine_status": "existing" if existing else "new",
        "existing_medicine_id": str(existing.id) if existing else None,
        "_parsed": parsed,
    }


def summarize(results):
    valid = [r for r in results if r["is_valid"]]
    new_names = {normalize_name(r["mapped_data"]["medicine_name"]) for r in valid if r["medicine_status"] == "new"}
    existing_ids = {r["existing_medicine_id"] for r in valid if r["medicine_status"] == "existing"}
    return {
        "total": len(results),
        "valid": len(valid),
        "error_count": len(results) - len(valid),
        "warning_count": sum(len(r["warnings"]) for r in results),
        "new_medicines": len(new_names),
        "existing_medicines": len(existing_ids),
        "new_batches": len(valid),
    }


def public_record(result):
    return {key: value for key, value in result.items() if not key.startswith("_")}
